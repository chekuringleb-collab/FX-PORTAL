import os, csv, io, requests, datetime, logging, math, hashlib, time
from functools import wraps
from flask import Flask, render_template, jsonify, make_response, redirect, url_for, request, flash, abort, session
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY') or os.urandom(32)
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'sqlite:///users.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_SECURE'] = os.environ.get('FLASK_ENV') == 'production'

ADMIN_SECRET = os.environ.get('ADMIN_SECRET', 'fx-admin-2026')

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Пожалуйста, войдите в систему.'

API_URL = "https://api.exchangerate-api.com/v4/latest/USD"
FRANKFURTER = "https://api.frankfurter.app"
REQUEST_TIMEOUT = 8
MIN_USERNAME_LEN, MAX_USERNAME_LEN, MIN_PASSWORD_LEN = 3, 32, 8
CACHE_TTL = 600  # 10 минут

# ── Простой кеш в памяти ──────────────────────────────────────────────────────
_cache = {}

def cache_get(key):
    if key in _cache:
        val, ts = _cache[key]
        if time.time() - ts < CACHE_TTL:
            return val
        del _cache[key]
    return None

def cache_set(key, val):
    _cache[key] = (val, time.time())

RU_NAMES = {
    "USD":"Доллар США","EUR":"Евро","RUB":"Российский рубль",
    "CNY":"Китайский юань","GBP":"Британский фунт","JPY":"Японская иена",
    "KZT":"Казахстанский тенге","BYN":"Белорусский рубль","TRY":"Турецкая лира",
    "AED":"Дирхам ОАЭ","AMD":"Армянский драм","GEL":"Грузинский лари",
    "CHF":"Швейцарский франк","CAD":"Канадский доллар","AUD":"Австралийский доллар",
    "SEK":"Шведская крона","NOK":"Норвежская крона","PLN":"Польский злотый",
}
CURRENCY_FLAGS = {
    "USD":"🇺🇸","EUR":"🇪🇺","RUB":"🇷🇺","CNY":"🇨🇳","GBP":"🇬🇧",
    "JPY":"🇯🇵","KZT":"🇰🇿","BYN":"🇧🇾","TRY":"🇹🇷","AED":"🇦🇪",
    "AMD":"🇦🇲","GEL":"🇬🇪","CHF":"🇨🇭","CAD":"🇨🇦","AUD":"🇦🇺",
    "SEK":"🇸🇪","NOK":"🇳🇴","PLN":"🇵🇱",
}
FRANKFURTER_SUPPORTED = {"EUR","GBP","JPY","CNY","CHF","CAD","AUD","SEK","NOK","PLN","TRY","AED","GEL"}
PERIODS = {"1w":7,"1m":30,"3m":90,"6m":180,"1y":365}

def validate_registration(username, password):
    errors = []
    if not username or len(username) < MIN_USERNAME_LEN:
        errors.append(f"Логин не короче {MIN_USERNAME_LEN} символов.")
    if len(username) > MAX_USERNAME_LEN:
        errors.append(f"Логин не длиннее {MAX_USERNAME_LEN} символов.")
    if not all(c.isalnum() or c in '_-' for c in username):
        errors.append("Логин: только буквы, цифры, _ и -.")
    if not password or len(password) < MIN_PASSWORD_LEN:
        errors.append(f"Пароль не короче {MIN_PASSWORD_LEN} символов.")
    return errors

# ── Models ────────────────────────────────────────────────────────────────────
class User(UserMixin, db.Model):
    __tablename__ = 'users'
    id          = db.Column(db.Integer, primary_key=True)
    username    = db.Column(db.String(MAX_USERNAME_LEN), unique=True, nullable=False, index=True)
    password    = db.Column(db.String(256), nullable=False)
    role        = db.Column(db.String(20), default='user', nullable=False)
    is_blocked  = db.Column(db.Boolean, default=False, nullable=False)
    theme       = db.Column(db.String(10), default='dark', nullable=False)
    favorites   = db.Column(db.String(200), default='', nullable=False)
    created_at  = db.Column(db.DateTime, default=datetime.datetime.utcnow)
    last_login  = db.Column(db.DateTime, nullable=True)
    login_count = db.Column(db.Integer, default=0, nullable=False)
    logs        = db.relationship('ActionLog', backref='user', lazy=True, cascade='all, delete-orphan')

    def set_password(self, raw):
        self.password = generate_password_hash(raw, method='pbkdf2:sha256:600000')
    def check_password(self, raw):
        return check_password_hash(self.password, raw)
    @property
    def is_admin(self): return self.role == 'admin'
    @property
    def favorites_list(self):
        return [f for f in self.favorites.split(',') if f] if self.favorites else []

class ActionLog(db.Model):
    __tablename__ = 'action_logs'
    id         = db.Column(db.Integer, primary_key=True)
    user_id    = db.Column(db.Integer, db.ForeignKey('users.id'), nullable=False)
    action     = db.Column(db.String(200), nullable=False)
    ip_address = db.Column(db.String(50), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)

def log_action(user_id, action):
    try:
        entry = ActionLog(user_id=user_id, action=action, ip_address=request.remote_addr)
        db.session.add(entry); db.session.commit()
    except Exception as e:
        logger.warning("Log error: %s", e)

@login_manager.user_loader
def load_user(uid): return db.session.get(User, int(uid))

def admin_required(f):
    @wraps(f)
    def dec(*a, **kw):
        if not current_user.is_authenticated or not current_user.is_admin: abort(403)
        return f(*a, **kw)
    return dec

# ── Хелпер получения курсов с кешем ──────────────────────────────────────────
def fetch_rates():
    cached = cache_get('rates')
    if cached: return cached
    resp = requests.get(API_URL, timeout=REQUEST_TIMEOUT)
    resp.raise_for_status()
    data = resp.json()
    cache_set('rates', data)
    return data

# ── Auth ──────────────────────────────────────────────────────────────────────
@app.route('/login', methods=['GET','POST'])
def login():
    if current_user.is_authenticated: return redirect(url_for('index'))
    if request.method == 'POST':
        username = request.form.get('username','').strip()
        password = request.form.get('password','')
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            if user.is_blocked:
                flash('Ваш аккаунт заблокирован.', 'error')
                return render_template('login.html')
            login_user(user, remember=bool(request.form.get('remember')))
            user.last_login = datetime.datetime.utcnow()
            user.login_count = (user.login_count or 0) + 1
            db.session.commit()
            log_action(user.id, "Вход в систему")
            nxt = request.args.get('next')
            return redirect(nxt if nxt and nxt.startswith('/') else url_for('index'))
        flash('Неверный логин или пароль.', 'error')
    return render_template('login.html')

@app.route('/register', methods=['GET','POST'])
def register():
    if current_user.is_authenticated: return redirect(url_for('index'))
    if request.method == 'POST':
        username = request.form.get('username','').strip()
        password = request.form.get('password','')
        confirm  = request.form.get('confirm_password','')
        errors = validate_registration(username, password)
        if password != confirm: errors.append("Пароли не совпадают.")
        if User.query.filter_by(username=username).first(): errors.append("Пользователь с таким логином уже существует.")
        if errors:
            for e in errors: flash(e, 'error')
            return render_template('register.html')
        user = User(username=username, role='user')
        user.set_password(password)
        db.session.add(user); db.session.commit()
        log_action(user.id, "Регистрация аккаунта")
        flash('Аккаунт создан! Войдите в систему.', 'success')
        return redirect(url_for('login'))
    return render_template('register.html')

@app.route('/admin-register/<secret>', methods=['GET','POST'])
def admin_register(secret):
    if secret != ADMIN_SECRET: abort(404)
    if current_user.is_authenticated: return redirect(url_for('index'))
    if request.method == 'POST':
        username = request.form.get('username','').strip()
        password = request.form.get('password','')
        confirm  = request.form.get('confirm_password','')
        errors = validate_registration(username, password)
        if password != confirm: errors.append("Пароли не совпадают.")
        if User.query.filter_by(username=username).first(): errors.append("Пользователь с таким логином уже существует.")
        if errors:
            for e in errors: flash(e, 'error')
            return render_template('admin_register.html', secret=secret)
        user = User(username=username, role='admin')
        user.set_password(password)
        db.session.add(user); db.session.commit()
        log_action(user.id, "Регистрация admin-аккаунта")
        flash('Администратор создан!', 'success')
        return redirect(url_for('login'))
    return render_template('admin_register.html', secret=secret)

@app.route('/logout', methods=['POST'])
@login_required
def logout():
    log_action(current_user.id, "Выход из системы")
    logout_user(); return redirect(url_for('login'))

# ── Pages ─────────────────────────────────────────────────────────────────────
@app.route('/')
@login_required
def index():
    return render_template('index.html', user=current_user, favorites=current_user.favorites_list)

@app.route('/history/<code>')
@login_required
def history_page(code):
    code = code.upper()
    if code not in RU_NAMES: abort(404)
    log_action(current_user.id, f"Открыл историю курса {code}")
    return render_template('history.html', currency_code=code, currency_name=RU_NAMES.get(code, code))

@app.route('/compare')
@login_required
def compare_page():
    return render_template('compare.html')

@app.route('/upload')
@login_required
def upload_page():
    return render_template('upload.html')

@app.route('/calculator')
@login_required
def calculator_page():
    return render_template('calculator.html')

# ── Профиль ───────────────────────────────────────────────────────────────────
@app.route('/profile', methods=['GET','POST'])
@login_required
def profile():
    if request.method == 'POST':
        action = request.form.get('action')

        if action == 'change_password':
            old_pw  = request.form.get('old_password','')
            new_pw  = request.form.get('new_password','')
            confirm = request.form.get('confirm_password','')
            if not current_user.check_password(old_pw):
                flash('Неверный текущий пароль.', 'error')
            elif len(new_pw) < MIN_PASSWORD_LEN:
                flash(f'Новый пароль не короче {MIN_PASSWORD_LEN} символов.', 'error')
            elif new_pw != confirm:
                flash('Пароли не совпадают.', 'error')
            else:
                current_user.set_password(new_pw)
                db.session.commit()
                log_action(current_user.id, "Сменил пароль")
                flash('Пароль успешно изменён!', 'success')

        elif action == 'change_theme':
            theme = request.form.get('theme', 'dark')
            if theme in ('dark', 'light'):
                current_user.theme = theme
                db.session.commit()
                flash('Тема изменена!', 'success')

        return redirect(url_for('profile'))

    logs = ActionLog.query.filter_by(user_id=current_user.id).order_by(ActionLog.created_at.desc()).limit(10).all()
    return render_template('profile.html', user=current_user, logs=logs)

# ── Избранное ─────────────────────────────────────────────────────────────────
@app.route('/api/favorites/toggle/<code>', methods=['POST'])
@login_required
def toggle_favorite(code):
    code = code.upper()
    if code not in RU_NAMES:
        return jsonify({"error": "Неизвестная валюта"}), 404
    favs = current_user.favorites_list
    if code in favs:
        favs.remove(code)
        action = "removed"
    else:
        favs.append(code)
        action = "added"
    current_user.favorites = ','.join(favs)
    db.session.commit()
    return jsonify({"action": action, "favorites": favs})

# ── Admin ─────────────────────────────────────────────────────────────────────
@app.route('/admin')
@login_required
@admin_required
def admin_panel():
    users = User.query.order_by(User.created_at.desc()).all()
    total_logins = db.session.query(db.func.sum(User.login_count)).scalar() or 0
    blocked_count = User.query.filter_by(is_blocked=True).count()
    return render_template('admin.html', users=users, total_logins=total_logins, blocked_count=blocked_count)

@app.route('/admin/delete/<int:uid>', methods=['POST'])
@login_required
@admin_required
def delete_user(uid):
    if uid == current_user.id:
        flash('Нельзя удалить себя.', 'error'); return redirect(url_for('admin_panel'))
    user = db.session.get(User, uid)
    if not user: abort(404)
    log_action(current_user.id, f"Удалил пользователя {user.username}")
    db.session.delete(user); db.session.commit()
    flash(f'Пользователь {user.username} удалён.', 'success')
    return redirect(url_for('admin_panel'))

@app.route('/admin/toggle-role/<int:uid>', methods=['POST'])
@login_required
@admin_required
def toggle_role(uid):
    if uid == current_user.id:
        flash('Нельзя изменить свою роль.', 'error'); return redirect(url_for('admin_panel'))
    user = db.session.get(User, uid)
    if not user: abort(404)
    old_role = user.role
    user.role = 'user' if user.role == 'admin' else 'admin'
    db.session.commit()
    log_action(current_user.id, f"Изменил роль {user.username}: {old_role} → {user.role}")
    flash(f'Роль {user.username} изменена на {user.role}.', 'success')
    return redirect(url_for('admin_panel'))

@app.route('/admin/toggle-block/<int:uid>', methods=['POST'])
@login_required
@admin_required
def toggle_block(uid):
    if uid == current_user.id:
        flash('Нельзя заблокировать себя.', 'error'); return redirect(url_for('admin_panel'))
    user = db.session.get(User, uid)
    if not user: abort(404)
    user.is_blocked = not user.is_blocked
    db.session.commit()
    action = "заблокировал" if user.is_blocked else "разблокировал"
    log_action(current_user.id, f"{action.capitalize()} пользователя {user.username}")
    flash(f'Пользователь {user.username} {"заблокирован" if user.is_blocked else "разблокирован"}.', 'success')
    return redirect(url_for('admin_panel'))

@app.route('/admin/logs')
@login_required
@admin_required
def admin_logs():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('q', '').strip()
    query = ActionLog.query
    if search:
        query = query.join(User).filter(
            db.or_(User.username.ilike(f'%{search}%'), ActionLog.action.ilike(f'%{search}%'))
        )
    logs = query.order_by(ActionLog.created_at.desc()).paginate(page=page, per_page=50, error_out=False)
    return render_template('admin_logs.html', logs=logs, search=search)

# ── API ───────────────────────────────────────────────────────────────────────
@app.route('/api/rates')
@login_required
def get_rates():
    try:
        data = fetch_rates()
        enriched = [
            {"code":c,"rate":round(r,4),"name":RU_NAMES.get(c,c),"flag":CURRENCY_FLAGS.get(c,"🏳️")}
            for c,r in data.get('rates',{}).items() if c in RU_NAMES
        ]
        enriched.sort(key=lambda x: x['code'])
        return jsonify({"rates":enriched,"base":"USD","date":data.get("date"),"cached": cache_get('rates') is not None})
    except requests.Timeout:
        return jsonify({"error":"Сервис временно недоступен"}), 503
    except requests.RequestException as e:
        logger.error("rates error: %s", e)
        return jsonify({"error":"Ошибка получения данных"}), 502

@app.route('/api/top-movers')
@login_required
def api_top_movers():
    """Топ рост и падение за сегодня (сравниваем с вчера через Frankfurter)"""
    try:
        data = fetch_rates()
        rates_today = data.get('rates', {})
        today = datetime.date.today()
        yesterday = today - datetime.timedelta(days=1)
        resp = requests.get(f"{FRANKFURTER}/{yesterday}?from=USD", timeout=REQUEST_TIMEOUT)
        resp.raise_for_status()
        rates_yesterday = resp.json().get('rates', {})
        movers = []
        for code in RU_NAMES:
            if code == 'USD': continue
            t = rates_today.get(code)
            y = rates_yesterday.get(code)
            if t and y and y != 0:
                change = round(((t - y) / y) * 100, 2)
                movers.append({"code": code, "name": RU_NAMES[code], "flag": CURRENCY_FLAGS.get(code,"🏳️"), "rate": round(t,4), "change": change})
        movers.sort(key=lambda x: x['change'], reverse=True)
        return jsonify({"top": movers[:3], "bottom": movers[-3:][::-1]})
    except Exception as e:
        logger.warning("top-movers error: %s", e)
        return jsonify({"top": [], "bottom": []}), 200

@app.route('/api/history/<code>')
@login_required
def api_history(code):
    code = code.upper()
    if code not in RU_NAMES: return jsonify({"error":"Неизвестная валюта"}), 404
    period = request.args.get('period','1m')
    days = PERIODS.get(period, 30)
    today = datetime.date.today()
    start = today - datetime.timedelta(days=days)
    cache_key = f"history_{code}_{period}"
    cached = cache_get(cache_key)
    if cached: return jsonify(cached)
    if code in FRANKFURTER_SUPPORTED:
        try:
            url = f"{FRANKFURTER}/{start}..{today}?from=USD&to={code}"
            resp = requests.get(url, timeout=REQUEST_TIMEOUT); resp.raise_for_status()
            fdata = resp.json()
            rates_dict = fdata.get('rates', {})
            labels = sorted(rates_dict.keys())
            if len(labels) > 30:
                step = len(labels) // 30
                labels = labels[::step]
            rates_list = [round(rates_dict[d][code], 4) for d in labels if code in rates_dict.get(d,{})]
            labels_fmt = [datetime.date.fromisoformat(d).strftime('%d.%m') for d in labels if code in rates_dict.get(d,{})]
            result = {"labels":labels_fmt,"rates":rates_list,"source":"ECB/Frankfurter","real":True}
            cache_set(cache_key, result)
            return jsonify(result)
        except requests.RequestException as e:
            logger.warning("Frankfurter fallback for %s: %s", code, e)
    try:
        data = fetch_rates()
        current_rate = data.get('rates', {}).get(code, 1.0)
    except requests.RequestException:
        current_rate = 1.0
    seed = int(hashlib.md5(code.encode()).hexdigest(), 16) % 1000
    labels, rates_list = [], []
    points = min(days, 30)
    for i in range(points, -1, -1):
        d = today - datetime.timedelta(days=int(days * i / points))
        labels.append(d.strftime('%d.%m'))
        variation = math.sin((seed + i) * 0.7) * 0.04
        rates_list.append(round(current_rate * (1 + variation), 4))
    result = {"labels":labels,"rates":rates_list,"source":"Симуляция","real":False}
    cache_set(cache_key, result)
    return jsonify(result)

@app.route('/api/admin/stats')
@login_required
@admin_required
def api_admin_stats():
    users = User.query.all()
    return jsonify({
        "total_users": len(users),
        "admins": sum(1 for u in users if u.is_admin),
        "blocked": sum(1 for u in users if u.is_blocked),
        "total_logins": sum(u.login_count or 0 for u in users),
    })

# ── Export ────────────────────────────────────────────────────────────────────
@app.route('/download')
@login_required
def download_csv():
    try:
        data = fetch_rates()
        rates = data.get('rates', {})
    except requests.RequestException:
        flash("Не удалось загрузить данные.", "error"); return redirect(url_for('index'))
    log_action(current_user.id, "Скачал CSV")
    si = io.StringIO()
    w = csv.writer(si, delimiter=';')
    w.writerow(['Код','Валюта','Флаг','Курс к USD','Дата'])
    today = datetime.date.today().isoformat()
    for c, r in sorted(rates.items()):
        if c in RU_NAMES: w.writerow([c, RU_NAMES[c], CURRENCY_FLAGS.get(c,''), round(r,4), today])
    out = make_response(si.getvalue().encode('utf-8-sig'))
    out.headers["Content-Disposition"] = f"attachment; filename=rates_{today}.csv"
    out.headers["Content-type"] = "text/csv; charset=utf-8"
    return out

@app.route('/download_xlsx')
@login_required
def download_xlsx():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        flash("openpyxl не установлен.", "error"); return redirect(url_for('index'))
    try:
        data = fetch_rates()
        rates = data.get('rates', {})
    except requests.RequestException:
        flash("Не удалось загрузить данные.", "error"); return redirect(url_for('index'))
    log_action(current_user.id, "Скачал Excel")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Курсы валют"
    header_fill = PatternFill("solid", fgColor="1a2540")
    header_font = Font(bold=True, color="4f8ef7", size=11)
    for col, h in enumerate(['Код','Валюта','Флаг','Курс к USD','Дата'], 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    today = datetime.date.today().isoformat()
    for c, r in sorted(rates.items()):
        if c in RU_NAMES: ws.append([c, RU_NAMES[c], CURRENCY_FLAGS.get(c,''), round(r,4), today])
    ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 28; ws.column_dimensions['D'].width = 14
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    out = make_response(buf.read())
    out.headers["Content-Disposition"] = f"attachment; filename=rates_{today}.xlsx"
    out.headers["Content-type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return out

# ── Errors ────────────────────────────────────────────────────────────────────
@app.errorhandler(403)
def e403(e): return render_template('error.html', code=403, message="Доступ запрещён"), 403
@app.errorhandler(404)
def e404(e): return render_template('error.html', code=404, message="Страница не найдена"), 404
@app.errorhandler(500)
def e500(e): return render_template('error.html', code=500, message="Внутренняя ошибка сервера"), 500

with app.app_context():
    db.create_all()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=os.environ.get('FLASK_ENV') != 'production')
